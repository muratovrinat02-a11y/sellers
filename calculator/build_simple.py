"""Простой калькулятор-КП: один базовый сценарий, только выручка.

Запуск: python3 calculator/build_simple.py calculator/kp-prostoy.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1] if len(sys.argv) > 1 else "kp-prostoy.xlsx"
MONTHS = 24
F = "Arial"

INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
RESULT_FILL = PatternFill("solid", fgColor="E2EFDA")
SOFT_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
RUB = '#,##0 "₽"'
NUM = "#,##0"
PCT = "0%"
PCT1 = "0.0%"

wb = Workbook()
kp = wb.active
kp.title = "КП"


def font(size=10, bold=False, color="000000", italic=False):
    return Font(name=F, size=size, bold=bold, color=color, italic=italic)


def band(row, text, cols=3):
    kp.cell(row=row, column=1, value=text).font = font(11, True, "FFFFFF")
    for c in range(1, cols + 1):
        kp.cell(row=row, column=c).fill = HEAD_FILL


REF = {}


def field(row, label, value, fmt, key, note="", fill=INPUT_FILL, color="0000FF"):
    kp.cell(row=row, column=1, value=label).font = font()
    c = kp.cell(row=row, column=2, value=value)
    c.font = font(color=color)
    c.fill = fill
    c.border = BOX
    c.alignment = Alignment(horizontal="right")
    if fmt:
        c.number_format = fmt
    kp.cell(row=row, column=3, value=note).font = font(9, color="808080", italic=True)
    REF[key] = f"$B${row}"


# --- шапка
kp["A1"] = '="Свой интернет-магазин для "&B5'
kp["A1"].font = font(16, True)
kp["A2"] = "Расчёт на ваших цифрах"
kp["A2"].font = font(10, color="808080", italic=True)

# --- ввод
band(4, "Данные бренда")
field(5, "Бренд", "Пример бренда", None, "brand")
field(6, "Оборот, ₽/мес", 5000000, RUB, "rev")
field(7, "Средний чек, ₽", 2500, RUB, "aov")
field(8, "Запросы бренда в Яндексе, в мес", 10000, NUM, "wordstat", "wordstat.yandex.ru")
field(9, "Подписчики в соцсетях и Telegram", 20000, NUM, "followers")
field(10, "Покупок одного клиента в год", 3, "0.0", "freq", "Одежда 2–3, косметика 4–6")

# --- результат (заполняется ниже, после расчёта)
band(12, "Что даст свой интернет-магазин")

# --- допущения базового сценария (мелко, внизу)
A0 = 24
band(A0, "Допущения (базовый сценарий, можно не трогать)")
GREY = "595959"
assum = [
    ("Покупатели, перешедшие в свою базу (QR, гарантия, чат), % заказов", 0.03, PCT1, "conv_base"),
    ("Повторные покупки базы, сделанные на сайте", 0.35, PCT, "base_site"),
    ("Отток базы в месяц", 0.05, PCT, "churn"),
    ("Доля поиска бренда, которую забирает сайт", 0.2, PCT, "brand_share"),
    ("Конверсия поиска бренда в заказ", 0.03, PCT1, "brand_cr"),
    ("Подписчики, покупающие на сайте в месяц", 0.004, PCT1, "social_cr"),
    ("Средний чек на сайте к текущему", 1.1, PCT, "aov_k"),
    ("Месяцев до выхода на полную мощность", 6, "0", "ramp"),
    ("Комиссии и расходы маркетплейса, % выручки", 0.42, PCT, "mp_take"),
]
for i, (label, v, fmt, key) in enumerate(assum):
    field(A0 + 1 + i, label, v, fmt, key, fill=SOFT_FILL, color=GREY)
    kp.cell(row=A0 + 1 + i, column=1).font = font(9, color=GREY)
    kp.cell(row=A0 + 1 + i, column=2).font = font(9, color=GREY)

src = A0 + 1 + len(assum)
kp.cell(row=src, column=1,
        value="Комиссии маркетплейсов: 42% выручки по данным сервиса «Финансист» (RB.ru, 25.09.2026).").font = font(8, color="808080", italic=True)


def K(key):
    return f"КП!{REF[key]}"


# --- расчёт (отдельный лист, 24 месяца)
calc = wb.create_sheet("Расчёт")
calc["A1"] = "Помесячный расчёт, базовый сценарий"
calc["A1"].font = font(12, True)
rows = ["Месяц", "Разгон", "Новые в базе", "База", "Заказы: база", "Заказы: поиск бренда",
        "Заказы: соцсети", "Заказы всего", "Выручка магазина, ₽", "Нарастающим итогом, ₽"]
R = {name: 3 + i for i, name in enumerate(rows)}
for name, r in R.items():
    calc.cell(row=r, column=1, value=name).font = font(bold=name in ("Месяц", "Выручка магазина, ₽"))
for m in range(1, MONTHS + 1):
    col = 1 + m
    c, p = L(col), L(col - 1)

    def at(n):
        return f"{c}{R[n]}"

    def prev(n):
        return f"{p}{R[n]}"

    f = {
        "Месяц": m,
        "Разгон": f"=MIN(1,{at('Месяц')}/{K('ramp')})",
        "Новые в базе": f"={K('rev')}/{K('aov')}*{K('conv_base')}*{at('Разгон')}+{at('Заказы: поиск бренда')}+{at('Заказы: соцсети')}",
        "База": f"={at('Новые в базе')}" if m == 1 else f"={prev('База')}*(1-{K('churn')})+{at('Новые в базе')}",
        "Заказы: база": "=0" if m == 1 else f"={prev('База')}*{K('freq')}/12*{K('base_site')}",
        "Заказы: поиск бренда": f"={K('wordstat')}*{K('brand_share')}*{K('brand_cr')}*{at('Разгон')}",
        "Заказы: соцсети": f"={K('followers')}*{K('social_cr')}*{at('Разгон')}",
        "Заказы всего": f"={at('Заказы: база')}+{at('Заказы: поиск бренда')}+{at('Заказы: соцсети')}",
        "Выручка магазина, ₽": f"={at('Заказы всего')}*{K('aov')}*{K('aov_k')}",
        "Нарастающим итогом, ₽": f"={at('Выручка магазина, ₽')}" if m == 1 else f"={prev('Нарастающим итогом, ₽')}+{at('Выручка магазина, ₽')}",
    }
    fmts = {"Месяц": "0", "Разгон": PCT, "Выручка магазина, ₽": RUB, "Нарастающим итогом, ₽": RUB}
    for name, r in R.items():
        cell = calc.cell(row=r, column=col, value=f[name])
        cell.number_format = fmts.get(name, NUM)
        cell.font = font(bold=name in ("Месяц", "Выручка магазина, ₽"))
calc.column_dimensions["A"].width = 26
for m in range(MONTHS):
    calc.column_dimensions[L(2 + m)].width = 12
calc.freeze_panes = "B3"

c12, c24 = L(1 + 12), L(1 + 24)
REV = R["Выручка магазина, ₽"]
CUM = R["Нарастающим итогом, ₽"]

results = [
    ("Выручка магазина через год, ₽/мес", f"=Расчёт!{c12}{REV}", RUB, False),
    ("Выручка магазина через 2 года, ₽/мес", f"=Расчёт!{c24}{REV}", RUB, False),
    ("Прибавка к обороту через 2 года", f"=Расчёт!{c24}{REV}/{K('rev')}", PCT, True),
    ("Дополнительная выручка за 2 года, ₽", f"=Расчёт!{c24}{CUM}", RUB, True),
    ("Не уйдёт в комиссии маркетплейса за 2 года, ₽", f"=Расчёт!{c24}{CUM}*{K('mp_take')}", RUB, True),
    ("Покупателей в своей базе через 2 года", f"=Расчёт!{c24}{R['База']}", NUM, False),
]
for i, (label, f_, fmt, big) in enumerate(results):
    r = 13 + i
    a = kp.cell(row=r, column=1, value=label)
    b = kp.cell(row=r, column=2, value=f_)
    b.number_format = fmt
    b.alignment = Alignment(horizontal="right")
    a.font = font(12 if big else 10, big)
    b.font = font(13 if big else 10, big, "375623" if big else "000000")
    if big:
        a.fill = RESULT_FILL
        b.fill = RESULT_FILL

# почему сейчас
W = 13 + len(results) + 1
kp.cell(row=W, column=1, value="Почему сейчас").font = font(10, True)
why = [
    "Маркетплейсы удерживают уже около 42% выручки продавца — полтора года назад было около 20%.",
    "Летом 2026 пострадали склады с товаром больше 400 тыс. продавцов; полная компенсация — только до 500 тыс. ₽.",
    "Свои интернет-магазины в 2026 году растут на 25,8%, маркетплейсы — около 10% (INFOLine).",
]
for i, t in enumerate(why):
    kp.cell(row=W + 1 + i, column=1, value="• " + t).font = font(9, color="404040")

kp.column_dimensions["A"].width = 58
kp.column_dimensions["B"].width = 20
kp.column_dimensions["C"].width = 26
kp.sheet_view.showGridLines = False

wb.save(OUT)
print("saved", OUT)
