"""Собирает калькулятор КП для селлера (xlsx) — потом загружается в Google Sheets.

Запуск: python3 calculator/build_calculator.py calculator/kalkulyator-kp-sellera.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1] if len(sys.argv) > 1 else "kalkulyator-kp-sellera.xlsx"
MONTHS = 24

FONT = "Arial"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="E2EFDA")
WHITE_BOLD = Font(name=FONT, bold=True, color="FFFFFF")
BOLD = Font(name=FONT, bold=True)
BLUE = Font(name=FONT, color="0000FF")
NORMAL = Font(name=FONT)
GREY = Font(name=FONT, color="666666", italic=True, size=9)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

RUB = '#,##0 "₽"'
NUM = "#,##0"
PCT = "0%"
PCT1 = "0.0%"

wb = Workbook()


def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font == Font() or c.font.name != FONT:
                c.font = Font(name=FONT, bold=c.font.bold, color=c.font.color,
                              italic=c.font.italic, size=c.font.size)


def header(ws, row, text, ncols):
    ws.cell(row=row, column=1, value=text).font = WHITE_BOLD
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = HEAD_FILL


# ---------------------------------------------------------------- Ввод
inp = wb.active
inp.title = "Ввод"
inp["A1"] = "Калькулятор КП: свой интернет-магазин для селлера"
inp["A1"].font = Font(name=FONT, bold=True, size=14)
inp["A2"] = "Жёлтые ячейки с синим шрифтом — вводные, меняй их. Остальное считается само. На каждого селлера: Файл → Создать копию."
inp["A2"].font = GREY

NAMES = {}  # label -> absolute ref on Ввод


def inp_row(r, label, value, fmt, key, note=""):
    inp.cell(row=r, column=1, value=label).font = NORMAL
    c = inp.cell(row=r, column=2, value=value)
    c.font = BLUE
    c.fill = INPUT_FILL
    c.border = BOX
    if fmt:
        c.number_format = fmt
    inp.cell(row=r, column=3, value=note).font = GREY
    NAMES[key] = f"Ввод!$B${r}"


header(inp, 4, "1. Профиль селлера (из профилирования)", 5)
seller_rows = [
    ("Бренд", "Пример бренда", None, "brand", ""),
    ("Ниша", "Одежда", None, "niche", ""),
    ("Оборот на маркетплейсах, ₽/мес", 3000000, RUB, "mp_rev", "MPStats / Moneyplace"),
    ("Средний чек на маркетплейсе, ₽", 2500, RUB, "mp_aov", "MPStats / Moneyplace"),
    ("Удержания маркетплейса, % выручки", 0.42, PCT, "mp_take", "Комиссия + логистика + хранение + реклама + возвраты. Рынок: 40–50% (Финансист, RB.ru 25.09.2026)"),
    ("Себестоимость товара, % от цены", 0.35, PCT, "cogs", "Спроси у селлера или возьми типовое для ниши"),
    ("Запросы бренда в Вордстате, в мес", 3000, NUM, "wordstat", "wordstat.yandex.ru: название бренда, латиница + кириллица"),
    ("Подписчики в соцсетях и Telegram", 5000, NUM, "followers", "Все свои каналы бренда вместе"),
    ("Покупок одного клиента в год (в нише)", 2, "0.0", "freq", "Одежда 2–3, косметика 4–6, расходники 6+"),
    ("Товар на складах маркетплейсов, ₽", 6000000, RUB, "stock", "Остатки по себестоимости или по цене"),
]
for i, (label, val, fmt, key, note) in enumerate(seller_rows):
    inp_row(5 + i, label, val, fmt, key, note)

# scenario assumptions: B/C/D
r0 = 17
header(inp, r0, "2. Допущения по сценариям", 5)
for j, name in enumerate(["Осторожный", "Базовый", "Сильный"]):
    c = inp.cell(row=r0 + 1, column=2 + j, value=name)
    c.font = BOLD
    c.fill = SUB_FILL
    c.alignment = Alignment(horizontal="center")
inp.cell(row=r0 + 1, column=1, value="Показатель").font = BOLD
inp.cell(row=r0 + 1, column=1).fill = SUB_FILL
inp.cell(row=r0 + 1, column=5, value="Откуда цифра").font = BOLD
inp.cell(row=r0 + 1, column=5).fill = SUB_FILL

SCEN = {}  # key -> row
scen_rows = [
    ("Покупатели МП, перешедшие в базу (QR, Telegram, гарантия), % от заказов", (0.01, 0.03, 0.05), PCT1, "conv_base",
     "QR на гарантию/инструкцию в Telegram-боте, адрес сайта в чате WB (разрешено). По рынку 1–5% заказов (не проверено). Нельзя: «бонус за отзыв» (штраф WB 50 000 ₽), контакты во вкладышах Ozon"),
    ("Доля повторных покупок базы, сделанных на сайте", (0.2, 0.35, 0.45), PCT, "base_site",
     "Зависит от выгоды на сайте: бонусы, наборы, эксклюзивы"),
    ("Доля брендового поиска, которую забирает сайт", (0.1, 0.2, 0.3), PCT, "brand_share",
     "Остальное уходит на WB/Ozon. 7 из 10 проверяют товар на МП (Я.Маркет, 2025)"),
    ("Конверсия брендового трафика", (0.02, 0.03, 0.04), PCT1, "brand_cr",
     "Брендовый трафик 3–6%, холодный 0,8–2% (оценка)"),
    ("Покупают на сайте в месяц, % подписчиков", (0.002, 0.004, 0.007), PCT1, "social_cr", "Оценка"),
    ("Бюджет на Директ / VK, ₽/мес", (30000, 60000, 100000), RUB, "ads_budget", "Холодная реклама на старте обычно в минус, окупается позже"),
    ("Цена клика, ₽", (45, 45, 45), RUB, "cpc", "Директ на поиске 44 ₽ (click.ru, 2 кв. 2026), товарные 50–150 ₽"),
    ("Конверсия рекламного трафика", (0.008, 0.012, 0.018), PCT1, "ads_cr", "Холодный трафик нового сайта 0,8–2% (оценка)"),
    ("Средний чек на сайте, % от чека МП", (1.0, 1.1, 1.2), PCT, "aov_k",
     "На сайте чек выше за счёт наборов и бесплатной доставки от суммы"),
    ("Новые покупатели (не ушли бы на МП), % выручки сайта", (0.2, 0.3, 0.4), PCT, "incr",
     "Остальное — каннибализация: эти деньги селлер получил бы на МП"),
    ("Отток базы в месяц (перестают покупать)", (0.06, 0.05, 0.04), PCT, "churn", "Без оттока база растёт бесконечно; 4–6% в месяц — около половины за год"),
    ("Месяцев до выхода на полную мощность", (9, 6, 4), "0", "ramp", "Разгон: SEO, база, оптимизация рекламы"),
]
for i, (label, vals, fmt, key, note) in enumerate(scen_rows):
    r = r0 + 2 + i
    inp.cell(row=r, column=1, value=label).font = NORMAL
    for j, v in enumerate(vals):
        c = inp.cell(row=r, column=2 + j, value=v)
        c.font = BLUE
        c.fill = INPUT_FILL
        c.border = BOX
        c.number_format = fmt
    inp.cell(row=r, column=5, value=note).font = GREY
    SCEN[key] = r

r1 = r0 + 2 + len(scen_rows) + 1
header(inp, r1, "3. Расходы своего сайта", 5)
cost_rows = [
    ("Эквайринг, % от выручки сайта", 0.02, PCT1, "acq", "СБП 0,4–0,7%, карты около 2,8% (ЮKassa)"),
    ("Доставка до ПВЗ, ₽/заказ (за счёт продавца)", 250, RUB, "ship", "Яндекс от 99 ₽, Ozon от 123 ₽, СДЭК от 250 ₽"),
    ("Сборка и упаковка, ₽/заказ", 70, RUB, "pick", "Фулфилмент 50–75 ₽"),
    ("Бонусы и рассылки, % от выручки базы", 0.07, PCT, "crm_pct", "Бонусы за покупку на сайте + сервис рассылок"),
    ("Брендовая реклама и SEO, % от выручки бренд-трафика", 0.05, PCT, "brand_cost", "Защита бренда в Директе"),
    ("Платформа, CRM, сервисы, ₽/мес", 15000, RUB, "platform", "InSales/Битрикс + CRM + бот"),
]
for i, (label, val, fmt, key, note) in enumerate(cost_rows):
    inp_row(r1 + 1 + i, label, val, fmt, key, note)

r2 = r1 + 1 + len(cost_rows) + 1
header(inp, r2, "4. Наши услуги (для КП)", 5)
svc_rows = [
    ("Запуск под ключ, ₽ (разово)", 200000, RUB, "svc_setup", "Подставь свою цену"),
    ("Сопровождение, ₽/мес", 50000, RUB, "svc_month", "Подставь свою цену"),
    ("Процент с продаж сайта", 0, PCT, "svc_pct", "Если работаешь за результат"),
]
for i, (label, val, fmt, key, note) in enumerate(svc_rows):
    inp_row(r2 + 1 + i, label, val, fmt, key, note)

inp.column_dimensions["A"].width = 62
for col in "BCD":
    inp.column_dimensions[col].width = 16
inp.column_dimensions["E"].width = 70
inp.column_dimensions["C"].width = 16
# column C in sections 1/3/4 holds notes; keep it readable via wrap off
inp.freeze_panes = "A4"

# ---------------------------------------------------------------- Расчёт
calc = wb.create_sheet("Расчёт")
calc["A1"] = "Помесячный расчёт по трём сценариям (24 месяца)"
calc["A1"].font = Font(name=FONT, bold=True, size=13)
calc["A2"] = "Ничего не вводить: всё считается из листа «Ввод»."
calc["A2"].font = GREY

FIRST = 2  # column B = month 1
LASTC = FIRST + MONTHS - 1

BLOCK_ROWS = [
    # key, label, fmt, formula(m_col, prev_col, scen_col) -> str
    ("month", "Месяц", "0", None),
    ("ramp", "Разгон (доля полной мощности)", PCT, None),
    ("mp_orders", "Заказов на МП в месяц", NUM, None),
    ("base_new", "Новых людей в базе (с МП + новые покупатели сайта)", NUM, None),
    ("base", "База на конец месяца", NUM, None),
    ("o_base", "Заказы сайта: повторные из базы", "#,##0.0", None),
    ("o_brand", "Заказы сайта: брендовый поиск", "#,##0.0", None),
    ("o_social", "Заказы сайта: соцсети и Telegram", "#,##0.0", None),
    ("o_ads", "Заказы сайта: Директ / VK", "#,##0.0", None),
    ("orders", "Заказов на сайте всего", NUM, None),
    ("aov", "Средний чек сайта, ₽", RUB, None),
    ("rev", "Выручка сайта, ₽", RUB, None),
    ("rev_incr", "  в т.ч. новые деньги (прирост оборота), ₽", RUB, None),
    ("rev_cann", "  в т.ч. перешло с МП, ₽", RUB, None),
    ("share", "Доля сайта в общем обороте", PCT1, None),
    ("costs", "Расходы сайта (без себестоимости), ₽", RUB, None),
    ("costs_pct", "Расходы сайта, % выручки", PCT, None),
    ("svc", "Наши услуги, ₽", RUB, None),
    ("profit_site", "Маржа сайта после себестоимости и расходов, ₽", RUB, None),
    ("profit_lost", "Маржа, которую эти заказы дали бы на МП, ₽", RUB, None),
    ("effect", "Чистый эффект для селлера, ₽", RUB, None),
    ("cum", "Нарастающим итогом, ₽", RUB, None),
    ("paid", "Окупилось (1 = да)", "0", None),
]
BLOCK_H = len(BLOCK_ROWS) + 2
ROWS = {}  # (scen_idx, key) -> row


def S(scen_idx, key):
    return f"{key}_{scen_idx + 1}"


def N(key):
    return key


# именованные диапазоны: короче формулы и понятнее при клике на ячейку
from openpyxl.workbook.defined_name import DefinedName
for key, ref in NAMES.items():
    wb.defined_names[key] = DefinedName(key, attr_text=ref)
for key, r in SCEN.items():
    for j in range(3):
        nm = f"{key}_{j + 1}"
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"Ввод!${L(2 + j)}${r}")


start = 4
for s, sname in enumerate(["Осторожный", "Базовый", "Сильный"]):
    top = start + s * BLOCK_H
    header(calc, top, f"Сценарий: {sname}", LASTC)
    for i, (key, label, fmt, _) in enumerate(BLOCK_ROWS):
        ROWS[(s, key)] = top + 1 + i
    for i, (key, label, fmt, _) in enumerate(BLOCK_ROWS):
        r = ROWS[(s, key)]
        lc = calc.cell(row=r, column=1, value=label)
        lc.font = BOLD if key in ("rev", "effect", "share") else NORMAL
        for m in range(1, MONTHS + 1):
            col = FIRST + m - 1
            c_ = L(col)
            p_ = L(col - 1)

            def R(k):
                return f"{c_}{ROWS[(s, k)]}"

            def P(k):
                return f"{p_}{ROWS[(s, k)]}"

            if key == "month":
                f = m
            elif key == "ramp":
                f = f"=MIN(1,{R('month')}/{S(s,'ramp')})"
            elif key == "mp_orders":
                f = f"={N('mp_rev')}/{N('mp_aov')}"
            elif key == "base_new":
                f = f"={R('mp_orders')}*{S(s,'conv_base')}*{R('ramp')}+{R('o_brand')}+{R('o_social')}+{R('o_ads')}"
            elif key == "base":
                f = f"={R('base_new')}" if m == 1 else f"={P('base')}*(1-{S(s,'churn')})+{R('base_new')}"
            elif key == "o_base":
                prev = "0" if m == 1 else P("base")
                f = f"={prev}*{N('freq')}/12*{S(s,'base_site')}"
            elif key == "o_brand":
                f = f"={N('wordstat')}*{S(s,'brand_share')}*{R('ramp')}*{S(s,'brand_cr')}"
            elif key == "o_social":
                f = f"={N('followers')}*{S(s,'social_cr')}*{R('ramp')}"
            elif key == "o_ads":
                f = f"={S(s,'ads_budget')}/{S(s,'cpc')}*{S(s,'ads_cr')}*(0.5+0.5*{R('ramp')})"
            elif key == "orders":
                f = f"={R('o_base')}+{R('o_brand')}+{R('o_social')}+{R('o_ads')}"
            elif key == "aov":
                f = f"={N('mp_aov')}*{S(s,'aov_k')}"
            elif key == "rev":
                f = f"={R('orders')}*{R('aov')}"
            elif key == "rev_incr":
                f = f"={R('rev')}*{S(s,'incr')}"
            elif key == "rev_cann":
                f = f"={R('rev')}-{R('rev_incr')}"
            elif key == "share":
                f = f"={R('rev')}/({N('mp_rev')}-{R('rev_cann')}+{R('rev')})"
            elif key == "costs":
                f = (f"={R('rev')}*{N('acq')}+{R('orders')}*({N('ship')}+{N('pick')})"
                     f"+{R('o_base')}*{R('aov')}*{N('crm_pct')}"
                     f"+{R('o_brand')}*{R('aov')}*{N('brand_cost')}"
                     f"+{S(s,'ads_budget')}+{N('platform')}")
            elif key == "costs_pct":
                f = f"=IF({R('rev')}=0,0,{R('costs')}/{R('rev')})"
            elif key == "svc":
                setup = f"{N('svc_setup')}+" if m == 1 else ""
                f = f"={setup}{N('svc_month')}+{R('rev')}*{N('svc_pct')}"
            elif key == "profit_site":
                f = f"={R('rev')}*(1-{N('cogs')})-{R('costs')}-{R('svc')}"
            elif key == "profit_lost":
                f = f"={R('rev_cann')}*(1-{N('cogs')}-{N('mp_take')})"
            elif key == "effect":
                f = f"={R('profit_site')}-{R('profit_lost')}"
            elif key == "cum":
                f = f"={R('effect')}" if m == 1 else f"={P('cum')}+{R('effect')}"
            elif key == "paid":
                f = f"=IF({R('cum')}>0,1,0)"
            cell = calc.cell(row=r, column=col, value=f)
            cell.number_format = fmt
            cell.font = BOLD if key in ("rev", "effect", "share") else NORMAL
            if key == "month":
                cell.font = BOLD
                cell.fill = SUB_FILL
                cell.alignment = Alignment(horizontal="center")
        if key == "month":
            lc.fill = SUB_FILL

calc.column_dimensions["A"].width = 50
for m in range(MONTHS):
    calc.column_dimensions[L(FIRST + m)].width = 12
calc.freeze_panes = "B4"

# ---------------------------------------------------------------- КП
kp = wb.create_sheet("КП", 0)
kp["A1"] = "=\"Свой интернет-магазин для \"&Ввод!$B$5"
kp["A1"].font = Font(name=FONT, bold=True, size=15)
kp["A2"] = "=\"Расчёт на ваших цифрах. Ниша: \"&Ввод!$B$6"
kp["A2"].font = GREY

header(kp, 4, "Сколько вы сейчас отдаёте маркетплейсам", 4)
rows_now = [
    ("Оборот на маркетплейсах, ₽/мес", f"={N('mp_rev')}", RUB),
    ("Удерживает маркетплейс, ₽/мес", f"={N('mp_rev')}*{N('mp_take')}", RUB),
    ("Удерживает маркетплейс, ₽/год", f"={N('mp_rev')}*{N('mp_take')}*12", RUB),
    ("Товар на складах маркетплейсов, ₽", f"={N('stock')}", RUB),
    ("Ваш бренд ищут в Яндексе, раз в мес", f"={N('wordstat')}", NUM),
]
for i, (label, f, fmt) in enumerate(rows_now):
    kp.cell(row=5 + i, column=1, value=label).font = NORMAL
    c = kp.cell(row=5 + i, column=2, value=f)
    c.number_format = fmt
    c.font = BOLD

top = 11
header(kp, top, "Что даст свой магазин", 4)
for j, name in enumerate(["Осторожный", "Базовый", "Сильный"]):
    c = kp.cell(row=top + 1, column=2 + j, value=name)
    c.font = BOLD
    c.fill = SUB_FILL
    c.alignment = Alignment(horizontal="center")
kp.cell(row=top + 1, column=1, value="Показатель").font = BOLD
kp.cell(row=top + 1, column=1).fill = SUB_FILL

c12 = L(FIRST + 11)
c24 = L(FIRST + 23)
y1 = (L(FIRST), L(FIRST + 11))
y2 = (L(FIRST + 12), L(FIRST + 23))


def cr(s, key, col):
    return f"Расчёт!{col}{ROWS[(s, key)]}"


def crange(s, key, a, b):
    return f"Расчёт!{a}{ROWS[(s, key)]}:{b}{ROWS[(s, key)]}"


kp_rows = [
    ("Выручка сайта на 12-й месяц, ₽/мес", lambda s: f"={cr(s,'rev',c12)}", RUB, False),
    ("Выручка сайта на 24-й месяц, ₽/мес", lambda s: f"={cr(s,'rev',c24)}", RUB, True),
    ("Доля сайта в обороте на 24-й месяц", lambda s: f"={cr(s,'share',c24)}", PCT1, True),
    ("Прирост общего оборота на 24-й месяц", lambda s: f"={cr(s,'rev_incr',c24)}/{N('mp_rev')}", PCT1, False),
    ("Выручка сайта за 2-й год, ₽", lambda s: f"=SUM({crange(s,'rev',*y2)})", RUB, False),
    ("Покупателей в своей базе через 24 мес", lambda s: f"={cr(s,'base',c24)}", NUM, False),
    ("Расходы сайта на 24-й месяц, % выручки", lambda s: f"={cr(s,'costs_pct',c24)}", PCT, False),
    ("Чистый эффект за 1-й год, ₽", lambda s: f"=SUM({crange(s,'effect',*y1)})", RUB, False),
    ("Чистый эффект за 2-й год, ₽", lambda s: f"=SUM({crange(s,'effect',*y2)})", RUB, True),
    ("Окупаемость вложений, мес", lambda s: f"=IFERROR(MATCH(1,{crange(s,'paid',L(FIRST),c24)},0),\"больше 24\")", "0", True),
]
for i, (label, fn, fmt, key) in enumerate(kp_rows):
    r = top + 2 + i
    kp.cell(row=r, column=1, value=label).font = BOLD if key else NORMAL
    for s in range(3):
        c = kp.cell(row=r, column=2 + s, value=fn(s))
        c.number_format = fmt
        c.font = BOLD if key else NORMAL
        c.alignment = Alignment(horizontal="right")
        if key:
            c.fill = KEY_FILL
    if key:
        kp.cell(row=r, column=1).fill = KEY_FILL

note_r = top + 2 + len(kp_rows) + 1
notes = [
    "Чистый эффект = маржа сайта после себестоимости, всех расходов и наших услуг минус маржа, которую перешедшие с МП заказы и так дали бы на маркетплейсе.",
    "Источники рыночных цифр — лист «Справка». Допущения сценариев — лист «Ввод», раздел 2.",
    "Маркетплейсы остаются главным каналом: сайт — второй канал со своей базой покупателей и страховка от изменения условий и потерь на складах.",
]
for i, t in enumerate(notes):
    kp.cell(row=note_r + i, column=1, value=t).font = GREY

kp.column_dimensions["A"].width = 46
for col in "BCD":
    kp.column_dimensions[col].width = 18

# ---------------------------------------------------------------- Справка
ref = wb.create_sheet("Справка")
ref["A1"] = "Рыночные цифры для разговора (проверены на 06.10.2026)"
ref["A1"].font = Font(name=FONT, bold=True, size=13)
heads = ["Цифра", "Значение", "Тип", "Источник"]
for j, h in enumerate(heads):
    c = ref.cell(row=3, column=1 + j, value=h)
    c.font = WHITE_BOLD
    c.fill = HEAD_FILL
facts = [
    ("Доля выручки, которую удерживают WB/Ozon/ЯМ", "20% → 42% за 1,5 года", "оценка", "Финансист, RB.ru 25.09.2026"),
    ("Рост комиссий 2023–2025", "WB +58%, Ozon +63%", "оценка", "АПЭТ/Mpstats, Коммерсант 19.02.2026"),
    ("Рентабельность селлеров", "15% → 5%, к декабрю около 3%", "оценка", "Market Papa, Коммерсант 26.08.2026"),
    ("Комиссия Ozon с 06.04.2026", "до 55%, одежда 43–48%", "факт", "Ведомости 03.02.2026"),
    ("Комиссия WB с 07.07.2026", "+5 п.п. склад WB, +6 FBS, +20 витрина", "факт", "Коммерсант 07.07.2026"),
    ("Срок выплат WB", "около месяца, досрочно −4%", "факт", "Коммерсант 02.04.2026"),
    ("Атакованные склады WB, лето 2026", "17–22% площадей", "оценка", "Forbes / mail.ru"),
    ("Первые выплаты WB за потерянный товар", "5–10% себестоимости", "факт", "vbr.ru 24.07.2026"),
    ("Полная компенсация WB", "только при ущербе до 500 тыс. ₽", "факт", "msk1.ru 18.09.2026"),
    ("Рост независимого e-com в 2026", "+25,8% против около +10% у МП", "факт (прогноз)", "INFOLine, Коммерсант 01.10.2026"),
    ("Селлеры WB, рассматривающие свой сайт", "41,9%", "факт по пересказам", "Forbes, опрос 1–3.08.2026"),
    ("Селлеры с работающим сайтом", "около 5%", "оценка", "ТПП, Ведомости 27.05.2026"),
    ("Доля сайта у продавцов, запустивших его", "у половины больше 30%", "вендор", "Т-Банк, 2023"),
    ("Маркетплейсы в выручке селлеров от 1 млн/мес", "89–100% у 64%", "факт", "Точка × Data Insight, 2025"),
    ("Возвраты: сайт против маркетплейса", "около 30% против около 70%", "кейс", "Rivernord, biz360 16.04.2025"),
    ("CRM-канал в выручке магазина", "8–27%", "кейсы вендора", "Mindbox"),
    ("Клик в Директе на поиске", "44 ₽", "факт", "click.ru, 2 кв. 2026"),
    ("Доставка до ПВЗ", "Яндекс от 99 ₽, Ozon от 123 ₽, СДЭК от 250 ₽", "факт", "New Retail 10.07.2026"),
    ("«Телодвижения»: выручка", "2021: 102 млн → 2024: 4,7 млрд → 2025: >7 млрд ₽", "факт", "Business Vector; Forbes 30 до 30, 05.2026"),
    ("«Телодвижения»: свой сайт", "запущен только в 03.2026, доля не раскрыта", "факт", "Sostav, блог бренда 11.03.2026"),
    ("Bungly (детская одежда)", "маркетплейсы 56% продаж, свой сайт + франшиза", "факт", "shoppers.media"),
    ("Штраф WB за «бонус за отзыв» во вкладыше", "50 000 ₽ за факт", "факт", "Перечень штрафов WB с 01.10.2026"),
    ("Адрес своего сайта в чате WB", "разрешён (Telegram и соцсети — нет, штраф 25 000 ₽)", "факт", "Перечень штрафов WB, п. 8"),
]
for i, row in enumerate(facts):
    for j, v in enumerate(row):
        ref.cell(row=4 + i, column=1 + j, value=v).font = NORMAL
ref.column_dimensions["A"].width = 48
ref.column_dimensions["B"].width = 42
ref.column_dimensions["C"].width = 18
ref.column_dimensions["D"].width = 40

for ws in wb.worksheets:
    style_all(ws)

wb.save(OUT)
print("saved", OUT)
